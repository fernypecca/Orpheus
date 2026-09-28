"""Document text extraction (Context.dev `bytes`/parse equivalent, local + polite).

Direct URLs to PDF / DOCX / XLSX / PPTX are fetched with a single polite GET
(honest UA, robots-respecting, 20 MB cap like the reference API) and their
text lands in the same `text` field HTML pages use — downstream LLM code does
not care about the source format.

Fail-open per document: any unsupported type or parse error yields None, and
the caller falls back to the normal browser pipeline (which fail-closes
honestly if the URL is really unreadable).
"""

from __future__ import annotations

import io
import re
from urllib.parse import urlparse

MAX_DOC_BYTES = 20_000_000

_DOC_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pptx": "pptx",
}

_WS = re.compile(r"\s+")


def _clean(text: str) -> str:
    return _WS.sub(" ", text or "").strip()


def detect_doc_format(url: str, content_type: str | None = None) -> str | None:
    """'pdf' | 'docx' | 'xlsx' | 'pptx' | None. Extension first, MIME fallback."""
    path = urlparse(url).path.lower()
    for ext, fmt in _DOC_EXTENSIONS.items():
        if path.endswith(ext):
            return fmt
    ct = (content_type or "").lower()
    if "pdf" in ct:
        return "pdf"
    if "officedocument.wordprocessingml" in ct:
        return "docx"
    if "officedocument.spreadsheetml" in ct:
        return "xlsx"
    if "officedocument.presentationml" in ct:
        return "pptx"
    return None


def _extract_pdf(data: bytes) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages[:100]:  # cap: no runaway 1000-page catalogs
        try:
            pages.append(_clean(page.extract_text() or ""))
        except Exception:
            continue
    text = "\n\n".join(p for p in pages if p)
    return text, {"format": "pdf", "pages": len(reader.pages)}


def _extract_docx(data: bytes) -> tuple[str, dict]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paras = [_clean(p.text) for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            paras.append(" | ".join(_clean(c.text) for c in row.cells))
    text = "\n\n".join(p for p in paras if p)
    return text, {"format": "docx", "paragraphs": len(paras)}


def _extract_xlsx(data: bytes) -> tuple[str, dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, sheets = [], 0
    for ws in wb.worksheets:
        sheets += 1
        if sheets > 20:
            break
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [_clean(str(c)) for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
            if len(rows) >= 500:
                break
        if rows:
            parts.append(f"# {ws.title}\n" + "\n".join(rows))
    return "\n\n".join(parts), {"format": "xlsx", "sheets": sheets}


def _extract_pptx(data: bytes) -> tuple[str, dict]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides):
        if i >= 100:
            break
        bits = []
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    t = _clean(shape.text or "")
                    if t:
                        bits.append(t)
                elif shape.has_table:
                    for row in shape.table.rows:
                        bits.append(" | ".join(_clean(c.text) for c in row.cells))
            except Exception:
                continue
        if bits:
            slides.append(f"## Slide {i + 1}\n" + "\n".join(bits))
    return "\n\n".join(slides), {"format": "pptx", "slides": len(prs.slides)}


_EXTRACTORS = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "xlsx": _extract_xlsx,
    "pptx": _extract_pptx,
}


def extract_document(data: bytes, fmt: str) -> tuple[str, dict] | None:
    """(text, info) or None on any failure. Never raises."""
    try:
        fn = _EXTRACTORS.get(fmt)
        if fn is None or not data:
            return None
        text, info = fn(data)
        if not text or not text.strip():
            return None
        return text.strip(), info
    except Exception:
        return None
